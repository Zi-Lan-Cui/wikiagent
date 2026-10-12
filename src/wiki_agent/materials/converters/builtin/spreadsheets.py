"""Office 格式 builtin 转换器：xlsx → Markdown 表格。

旧二进制格式（.xls/.doc/.ppt）不在支持表内——loader 的模态表
本就不收它们。
"""

from __future__ import annotations

from typing import Any

from wiki_agent.log import get_logger

from ...loader import RawFileProperties
from ..base import BaseConverter, ConvertedFile

logger = get_logger("CONVERTER")


class XlsxConverter(BaseConverter):
    """每个工作表一节：表名 + 首行表头的 Markdown 表格。

    已知限制：公式取打开时缓存的算值（data_only），从未在 Excel 中
    算过公式的文件，公式格读为空——这是 openpyxl 的边界不是解析 bug。
    """

    def accepts(self, raw_file: RawFileProperties) -> bool:
        return raw_file.ext == "xlsx"

    async def convert(self, raw_file: RawFileProperties) -> ConvertedFile:
        # 延迟 import：非表格用途时 openpyxl 不进加载路径
        from openpyxl import load_workbook

        book = load_workbook(raw_file.path, read_only=True, data_only=True)
        try:
            sections = [
                rendered
                for ws in book.worksheets
                if (rendered := _sheet_to_markdown(ws.title, ws.iter_rows(values_only=True)))
            ]
        finally:
            book.close()
        content = "\n\n".join(sections)
        logger.info("  xlsx 转写: %s → %d chars (%d sheets)", raw_file.name, len(content), len(sections))
        return ConvertedFile.from_raw(raw_file, content)


def _cell(value: Any) -> str:
    """单元格 → 表格安全文本：管道转义、换行压平、None 留空。"""
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ").replace("\r", "")


def _sheet_to_markdown(title: str, rows: Any) -> str:
    """行迭代器 → `### 表名` + md 表格；空表返回空串。"""
    lines = [f"### {title}", ""]
    header_done = False
    for row in rows:
        cells = [_cell(v) for v in row]
        if not any(cells):
            continue
        lines.append("| " + " | ".join(cells) + " |")
        if not header_done:
            lines.append("|" + "---|" * len(cells))
            header_done = True
    if not header_done:
        return ""
    return "\n".join(lines)
